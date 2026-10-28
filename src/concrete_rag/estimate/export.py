"""Export an estimate to Excel (live formulas) and PDF."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from concrete_rag.config import get_settings
from concrete_rag.normalize.currency import format_inr

HDR = PatternFill("solid", fgColor="1F4E78")
WARN = PatternFill("solid", fgColor="FCE4D6")


def _header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=ws.max_row, column=i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w


def to_excel(est: dict, path: Path | None = None) -> Path:
    path = path or get_settings().output_dir / f"estimate_{est['id']}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append([f"Bill of Quantities - {est['project']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Schedule: {est['sor']} | Location: {est['location']} | Generated {est['created']}"])
    ws.append([])
    _header(ws, ["Sl", "DSR item", "Description", "Unit", "Quantity", "Rate (₹)", "Amount (₹)", "Citation", "Rate check"],
            [5, 11, 70, 7, 11, 12, 15, 34, 12])
    first = ws.max_row + 1
    for l in est["boq"]:
        ws.append([l["sl"], l["item_code"], l["description"], l["unit"], l["quantity"], l["rate"], None,
                   l["citation"], "confirmed" if l["rate_confirmed"] else "VERIFY"])
        r = ws.max_row
        ws.cell(r, 7, f"=E{r}*F{r}")
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        if not l["rate_confirmed"]:
            for c in range(1, 10):
                ws.cell(r, c).fill = WARN
    last = ws.max_row
    ws.append([])
    ws.append(["", "", "Sub-total", "", "", "", f"=SUM(G{first}:G{last})"])
    st = ws.max_row
    ws.append(["", "", f"Contingency @ {est['summary']['contingency_percent']} %", "", "", "",
               f"=G{st}*{est['summary']['contingency_percent']}/100"])
    ws.append(["", "", "Grand total", "", "", "", f"=G{st}+G{st + 1}"])
    for r in range(st, st + 3):
        ws.cell(r, 3).font = ws.cell(r, 7).font = Font(bold=True)
    for r in range(first, st + 3):
        for c in (6, 7):
            ws.cell(r, c).number_format = "#,##,##0.00"

    q = wb.create_sheet("Quantities")
    _header(q, ["Element", "Type", "Grade", "Level", "Nos", "Concrete m3", "Formwork m2", "Steel kg", "Notes"],
            [10, 14, 8, 14, 6, 12, 12, 10, 60])
    for r in est["quantities"]["rows"]:
        q.append([r["id"], r["type"], r["grade"], r["level"], r["count"], r["concrete_m3"], r["formwork_m2"],
                  r["steel_kg"], "; ".join(r["notes"])])

    if est.get("cement_statement"):
        c = wb.create_sheet("Cement statement")
        _header(c, ["DSR item", "Quantity", "Unit", "Cement qtl/unit", "Cement (qtl)", "Coefficient confirmed"],
                [11, 12, 8, 15, 14, 18])
        for x in est["cement_statement"]:
            c.append([x["item_code"], x["quantity"], x["unit"], x["coefficient_qtl_per_unit"], None, x["confirmed"]])
            c.cell(c.max_row, 5, f"=B{c.max_row}*D{c.max_row}")

    ch = est.get("cost_heads") or {}
    if ch.get("heads"):                                   # CPWD DAR split
        h = wb.create_sheet("Cost heads")
        heads = list(ch["heads"])
        _header(h, ["DSR item", "Amount", *[k.replace("_", " ").title() for k in heads], "DAR analysis"],
                [11, 14, *[13] * len(heads), 46])
        for x in ch["lines"]:
            h.append([x["item_code"], x["amount"], *[x["split"].get(k, 0) for k in heads], x["citation"]])
        h.append(["Total", ch["analysed_amount"], *[ch["heads"][k] for k in heads]])
        h.append(["Share %", None, *[ch["share_percent"][k] for k in heads]])
        if ch.get("not_analysed_items"):
            h.append([f"No DAR analysis: {', '.join(ch['not_analysed_items'])} ({ch['not_analysed_amount']})"])
        h.append([])
        h.append([ch["basis"]])
    elif ch.get("detail"):                                # estimates saved before DAR was ingested
        h = wb.create_sheet("Cost heads")
        h.append(["Materials (basic cost)", ch["materials_basic_cost"]])
        h.append(["Labour + T&P + formwork + overheads + GST (remainder)", ch["labour_equipment_overheads_gst_formwork"]])
        h.append([])
        h.append([ch["basis"]])
        h.column_dimensions["A"].width = 60

    n = wb.create_sheet("Assumptions & refs")
    n.column_dimensions["A"].width = 140
    n.append(["ASSUMPTIONS"])
    for a in est["assumptions"]:
        n.append([a])
    n.append([])
    n.append(["WARNINGS"])
    for w in est["warnings"]:
        n.append([w])
    n.append([])
    n.append(["REFERENCES"])
    for r in est["references"]:
        n.append([f"{r['code']}, {r['item']}"])
    wb.save(path)
    return path


def to_pdf(est: dict, path: Path | None = None) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    path = path or get_settings().output_dir / f"estimate_{est['id']}.pdf"
    ss = getSampleStyleSheet()
    small = ss["BodyText"].clone("small", fontSize=7.5, leading=9)
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    story = [Paragraph(f"Bill of Quantities - {est['project']}", ss["Title"]),
             Paragraph(f"{est['sor']} | Location: {est['location']} | {est['created']}", ss["Normal"]), Spacer(1, 8)]
    data = [["Sl", "Item", "Description", "Unit", "Qty", "Rate (Rs)", "Amount (Rs)", "Check"]]
    for l in est["boq"]:
        data.append([l["sl"], l["item_code"], Paragraph(l["description"][:400], small), l["unit"], f"{l['quantity']:.2f}",
                     f"{l['rate']:,.2f}", f"{l['amount']:,.2f}", "ok" if l["rate_confirmed"] else "VERIFY"])
    s = est["summary"]
    for label, v in (("Sub-total", s["subtotal"]), (f"Contingency {s['contingency_percent']}%", s["contingency"]),
                     ("Grand total", s["grand_total"])):
        data.append(["", "", label, "", "", "", f"{v:,.2f}", ""])
    t = Table(data, colWidths=[22, 48, 420, 34, 52, 62, 76, 40], repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
             ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (4, 1), (6, -1), "RIGHT")]
    for i, l in enumerate(est["boq"], 1):
        if not l["rate_confirmed"]:
            style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FCE4D6")))
    t.setStyle(TableStyle(style))
    story += [t, Spacer(1, 10), Paragraph(f"Grand total: {format_inr(s['grand_total']).replace('₹', 'Rs')}", ss["Heading3"]),
              Paragraph(s["note"], small), Spacer(1, 6), Paragraph("<b>Assumptions</b>", ss["Normal"])]
    story += [Paragraph("- " + a, small) for a in est["assumptions"]]
    if est["warnings"]:
        story += [Spacer(1, 6), Paragraph("<b>Warnings</b>", ss["Normal"])] + [Paragraph("- " + w, small) for w in est["warnings"]]
    doc.build(story)
    return path
